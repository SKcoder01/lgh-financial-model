"""
colour_finance.py — finance-facing colour scheme for the C12 io workbook.

One rule for finance: GREEN = update this,  GREY = leave alone (model fills it).

  brown   row 2 (dates) and row 3 (markers)
  yellow  text labels / metadata (col A, i/o, field_name, index)
  green   actuals (hard numbers left of the i/o marker) + model inputs (i-rows
          in the i/o block) + input-page input values
  grey    outputs (o-rows in the i/o block)
  white   everything else (c-rows, calcs, empties, formulas in the actuals zone)

Boundary between actuals and model is derived from the row-3 'i/o' marker, so it
auto-adjusts if columns shift. Saves via save_clean (no empty-<v> corruption).
"""
import openpyxl, re
from openpyxl.styles import PatternFill, Font
import c12_dept as cd
import save_clean

BROWN  = PatternFill(fill_type='solid', fgColor='FFCC9966')
YELLOW = PatternFill(fill_type='solid', fgColor='FFFFFF00')
GREEN  = PatternFill(fill_type='solid', fgColor='FF00B050')
GREY   = PatternFill(fill_type='solid', fgColor='FFD3D3D3')
WHITE  = PatternFill(fill_type='solid', fgColor='FFFFFFFF')

INPUT_PAGES = {'c12_inputs', '.c12_inputs.hr'}

def _row_markers(ws):
    m = {}
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if isinstance(b, str) and b.strip().lower() in ('i', 'o', 'c'):
            m[r] = b.strip().lower()
    return m

def apply_finance_colours(wb, sheets=None):
    tally = {}
    for name in (sheets or wb.sheetnames):
        ws = wb[name]
        _, io_start = cd.find_section_boundaries(ws)
        markers = _row_markers(ws)
        is_input_page = name in INPUT_PAGES
        counts = {'brown':0,'yellow':0,'green':0,'grey':0,'white':0}
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c); v = cell.value
                if r in (2, 3):
                    fill, key = BROWN, 'brown'
                elif isinstance(v, str) and not v.startswith('='):
                    fill, key = YELLOW, 'yellow'
                else:
                    mk = markers.get(r)
                    if is_input_page:
                        if mk == 'i' and c >= 5 and v is not None: fill, key = GREEN, 'green'
                        elif mk == 'o':                            fill, key = GREY,  'grey'
                        else:                                      fill, key = WHITE, 'white'
                    elif io_start and c < io_start:
                        # actuals zone: hard number = an actual to maintain
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            fill, key = GREEN, 'green'
                        else:
                            fill, key = WHITE, 'white'
                    elif io_start and c >= io_start:
                        if   mk == 'i': fill, key = GREEN, 'green'
                        elif mk == 'o': fill, key = GREY,  'grey'
                        else:           fill, key = WHITE, 'white'
                    else:
                        fill, key = WHITE, 'white'
                cell.fill = fill
                counts[key] += 1
        tally[name] = counts
    return tally

def colour_file(in_path, out_path, sheets=None, verbose=True):
    wb = openpyxl.load_workbook(in_path)
    tally = apply_finance_colours(wb, sheets)
    n = save_clean.save_clean(wb, out_path)
    if verbose:
        print(f"{'sheet':16} {'green':>6} {'grey':>6} {'yellow':>7} {'brown':>6} {'white':>6}")
        for s, c in tally.items():
            print(f"{s:16} {c['green']:>6} {c['grey']:>6} {c['yellow']:>7} {c['brown']:>6} {c['white']:>6}")
        print(f"\nsaved {out_path} ({n} empty <v> stripped)")
    return tally
