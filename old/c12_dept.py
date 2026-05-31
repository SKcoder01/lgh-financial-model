"""
c12_dept.py — shared engine for the C12 department models (HM, D&S, Bi&F).

All three departments share the SAME calculation; they differ only in:
  - which sheet they live on
  - which .c12_*_hr sheet feeds the HR summary rows
  - the two department-tagged output names (total_lgh_*_revenue / total_c12_*_revenue),
    which are discovered from the sheet by pattern — no hardcoding.

Calc is PURE. calibrate() compares Python vs Excel and WRITES NOTHING.
Fee/service totals are summed GENERICALLY (every `i` row in each section, found by the
always-present total_fees_earned / total_service_revenue anchor rows), so renaming or
adding fee lines in Excel can never silently drop one.
"""

import openpyxl
from datetime import date, datetime

EXCEL_FILE = 'financial_model_excel_claude.xlsx'

# ---- column / row detection -------------------------------------------------

def find_section_boundaries(ws):
    actual_start = io_start = None
    for col in range(1, ws.max_column + 1):
        v = ws.cell(3, col).value
        if isinstance(v, str):
            if v.strip().lower() == 'actual' and actual_start is None:
                actual_start = col
            elif v.strip().lower() == 'i/o' and io_start is None:
                io_start = col
    return actual_start, io_start

def get_opening_col(ws, projections_from):
    _, io_start = find_section_boundaries(ws)
    if not io_start:
        return None
    best_col = best_date = None
    for col in range(1, io_start):
        v = ws.cell(2, col).value
        if not isinstance(v, (date, datetime)):
            continue
        d = v.date() if isinstance(v, datetime) else v
        if d < projections_from and (best_date is None or d > best_date):
            best_date, best_col = d, col
    return best_col

def get_period_cols(ws, from_date, to_date):
    _, io_start = find_section_boundaries(ws)
    if not io_start:
        return {}
    cols = {}
    for col in range(io_start, ws.max_column + 1):
        v = ws.cell(2, col).value
        if not isinstance(v, (date, datetime)):
            continue
        d = v.date() if isinstance(v, datetime) else v
        if from_date <= d <= to_date:
            q = (d.month - 1) // 3 + 1
            cols[f"Q{q}_{d.year}"] = col
    return dict(sorted(cols.items(), key=lambda x: x[1]))

def find_row_by_field(ws, field, field_col=3):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, field_col).value == field:
            return r
    return None

def discover_revenue_fields(ws):
    """Find the two department-tagged output rows by pattern."""
    rev_lgh = rev_c12 = None
    for r in range(1, ws.max_row + 1):
        f = ws.cell(r, 3).value
        if not isinstance(f, str):
            continue
        if f.startswith('total_lgh_') and f.endswith('_revenue'):
            rev_lgh = f
        elif f.startswith('total_c12_') and f.endswith('_revenue'):
            rev_c12 = f
    return rev_lgh, rev_c12

# ---- read helpers -----------------------------------------------------------

def read_sheet(ws, cols, io_col=2, field_col=3, index_col=4):
    """Rows keyed by (field_name, index). Flags duplicate keys under '_dupes'."""
    data, dupes = {}, []
    for r in range(1, ws.max_row + 1):
        io = ws.cell(r, io_col).value
        if not isinstance(io, str) or io.strip().lower() not in ('i', 'o', 'c'):
            continue
        field = ws.cell(r, field_col).value
        if not field:
            continue
        index = ws.cell(r, index_col).value or ''
        key = (field, index)
        if key in data:
            dupes.append((key, r))
        entry = {'io': io.strip().lower()}
        for label, col in cols.items():
            v = ws.cell(r, col).value
            entry[label] = v if isinstance(v, (int, float)) else 0.0
        data[key] = entry
    data['_dupes'] = dupes
    return data

def get(data, field, index='', period='Q1_2026', default=0.0):
    return data.get((field, index), {}).get(period, default)

def read_c12_inputs(wb):
    ws = wb['c12_inputs']
    inputs = {}
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value != 'i':
            continue
        field = ws.cell(r, 3).value
        val = ws.cell(r, 5).value
        if not field:
            continue
        if isinstance(val, (int, float)):
            inputs[field] = val
        elif isinstance(val, datetime):
            inputs[field] = val.date()
        elif isinstance(val, date):
            inputs[field] = val
    return inputs

def read_dept_inputs(wb_data, sheet, period_cols, opening_col):
    ws = wb_data[sheet]
    cols = {'opening': opening_col}; cols.update(period_cols)
    return read_sheet(ws, cols)

def read_hr_outputs(wb_data, hr_sheet, period_cols, opening_col):
    ws = wb_data[hr_sheet]
    cols = {'opening': opening_col}; cols.update(period_cols)
    data = read_sheet(ws, cols)
    keep = {'hr_grand_total', 'hr_reimbursable', 'hr_not_reimbursable', 'hr_only_c12'}
    return {k: v for k, v in data.items() if isinstance(k, tuple) and k[0] in keep}

def section_total(ws, top_row, bottom_row, col, io_col=2):
    total = 0.0
    for r in range(top_row + 1, bottom_row):
        if ws.cell(r, io_col).value == 'i':
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)):
                total += v
    return total

def read_section_totals(ws, period_cols):
    r_fe = find_row_by_field(ws, 'total_fees_earned')
    r_fp = find_row_by_field(ws, 'total_fees_paid')
    r_sr = find_row_by_field(ws, 'total_service_revenue')
    fees = {p: section_total(ws, 4, r_fe, c) for p, c in period_cols.items()}
    serv = {p: section_total(ws, r_fp, r_sr, c) for p, c in period_cols.items()}
    return {'total_fees_earned': fees, 'total_service_revenue': serv}

# ---- the calculation (identical for every department) -----------------------

def calc_dept(inputs, hr_data, c12_inputs, section_totals, calc_periods, rev_lgh, rev_c12):
    cb = c12_inputs['chargeback_pct_actual']      # required assumption — no fallback
    pp = c12_inputs['participation_pct_actual']   # required assumption — no fallback
    mult = 1 / (1 - cb)
    res = {}

    def store(f, i, p, v): res.setdefault((f, i), {})[p] = v
    def inp(f, i='', p='Q1_2026'): return get(inputs, f, i, p)
    def hr(f, p): return get(hr_data, f, '', p)
    def prior(f, p):
        idx = calc_periods.index(p)
        if idx == 0:
            return get(inputs, f, '', 'opening', 0.0)
        return res.get((f, ''), {}).get(calc_periods[idx - 1], 0.0)

    for p in calc_periods:
        # FEE REVENUE
        tfe = section_totals['total_fees_earned'][p]
        df, pfa = inp('deferred_fee', '', p), inp('paid_fee_accruals', '', p)
        caf = df + prior('cumm_accrued_fee', p) + pfa
        pcf = tfe + df; tfp = pfa + pcf
        store('total_fees_earned', '', p, tfe); store('cumm_accrued_fee', '', p, caf)
        store('paid_current_fees', '', p, pcf); store('total_fees_paid', '', p, tfp)

        # SERVICE REVENUE
        tsr = section_totals['total_service_revenue'][p]
        dsf, psa = inp('deferred_service_fee', '', p), inp('paid_service_accruals', '', p)
        casf = dsf + prior('cumm_accrued_service_fee', p) + psa
        pcsf = tsr + dsf; tsrp = psa + pcsf
        store('total_service_revenue', '', p, tsr); store('cumm_accrued_service_fee', '', p, casf)
        store('paid_current_service_fees', '', p, pcsf); store('total_service_revenue_paid', '', p, tsrp)

        # CHARGEBACKS
        reim = hr('hr_reimbursable', p); tcb = reim * mult
        store('chargeback_multiplier', '', p, mult); store('reimbursable_by_hotels', '', p, reim)
        store('total_chargeback', '', p, tcb)

        # LGH REVENUE TOTALS  (rev_lgh discovered per sheet)
        store(rev_lgh, '', p, tsrp + tfp + tcb)
        store('lgh_ar_cumm', '', p, caf + casf)

        # C12 PARTICIPATION  (rev_c12 discovered per sheet)
        pme, cue = inp('pre_mod_extra_pct', '', p), inp('catchup_extra_pct', '', p)
        cpa = (pp + pme) * (pcf + pcsf) + (pp + cue) * (pfa + psa)
        cbf = cb * tcb
        store('c12_participation_pct', '', p, pp); store('chargeback_pct', '', p, cb)
        store('c12_participation_amt', '', p, cpa); store('chargeback_fees', '', p, cbf)
        store(rev_c12, '', p, cbf + cpa)
        store('increase_accrued_participation', '', p, -(pp + pme) * (df + dsf))
        store('decrease_accrued_participation', '', p, -(pp + cue) * (pfa + psa))

        # HR passthrough
        for f in ('hr_grand_total', 'hr_reimbursable', 'hr_not_reimbursable', 'hr_only_c12'):
            store(f, '', p, hr(f, p))
    return res

# ---- calibrate: compare Python vs Excel, write nothing ----------------------

def calibrate(sheet, hr_sheet, excel_file=EXCEL_FILE, tol=0.5, verbose=True):
    wb      = openpyxl.load_workbook(excel_file)
    wb_data = openpyxl.load_workbook(excel_file, data_only=True)

    ci = read_c12_inputs(wb)
    pf, pt = ci.get('projections_from'), ci.get('projections_to')
    ws = wb_data[sheet]
    opening = get_opening_col(ws, pf)
    period_cols = get_period_cols(ws, pf, pt)
    cps = list(period_cols.keys())

    rev_lgh, rev_c12 = discover_revenue_fields(ws)
    inputs = read_dept_inputs(wb_data, sheet, period_cols, opening)
    hr     = read_hr_outputs(wb_data, hr_sheet, period_cols, opening)
    sect   = read_section_totals(ws, period_cols)
    results = calc_dept(inputs, hr, ci, sect, cps, rev_lgh, rev_c12)

    dupes = inputs.get('_dupes', [])
    diffs = []
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value != 'o':
            continue
        f = ws.cell(r, 3).value; i = ws.cell(r, 4).value or ''
        if not f or (f, i) not in results:
            continue
        for p, c in period_cols.items():
            e = ws.cell(r, c).value; e = e if isinstance(e, (int, float)) else 0.0
            py = results[(f, i)].get(p)
            if py is None:
                continue
            if abs(py - e) > tol:
                diffs.append((f, i, p, e, py, py - e))

    if verbose:
        print(f"[{sheet}] periods {cps} | revenue fields: {rev_lgh}, {rev_c12}")
        if dupes:
            print(f"  WARNING duplicate (field,index) keys: {dupes}")
        if diffs:
            print(f"  FAIL {len(diffs)} differences vs Excel:")
            for f, i, p, e, py, d in diffs[:30]:
                print(f"    {f:30} {i:10} {p:9} excel={e:14,.2f} py={py:14,.2f} d={d:14,.2f}")
        else:
            print("  OK zero differences vs Excel")
    return results, diffs
