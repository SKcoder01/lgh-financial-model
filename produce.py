"""
produce.py — produce-mode writer for the C12 model.

Calibrate mode computes and compares (writes nothing). Produce mode writes the
computed o-values into the o-cells, REPLACING their formulas with hard numbers,
then saves via save_clean (strips empty <v>, forces recalc-on-open).

Writer mirrors diff_o_rows exactly: same o-row scan, same (field,index) keys,
same period_cols. The only difference is cell.value = py instead of comparing.
Writes go to the FORMULA workbook (wb), not the data_only one.
"""
import openpyxl
import c12_dept as cd
import c12_oc
import save_clean

DEPTS = [('c12_hm', '.c12_hm_hr'), ('c12_d&s', '.c12_d&s_hr'), ('c12_bi&F', '.c12_bi&f_hr')]

def write_o_rows(ws, results, period_cols):
    """Write Python results into every o-cell on ws. Returns list of (field,index,period,old,new)."""
    written = []
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value != 'o':
            continue
        f = ws.cell(r, 3).value; i = ws.cell(r, 4).value or ''
        if not f or (f, i) not in results:
            continue
        for p, c in period_cols.items():
            py = results[(f, i)].get(p)
            if py is None:
                continue
            old = ws.cell(r, c).value
            ws.cell(r, c).value = py          # overwrite formula with hard number
            written.append((f, i, p, old, py))
    return written

def produce(excel_file, out_path, verbose=True):
    # data_only for computing (needs cached input values); formula wb for writing
    wb,  wbd, ci, pf, pt = cd.load(excel_file)
    dept_results = {}
    pcs_by_sheet = {}
    for sheet, hr in DEPTS:
        res, pcs, rl, rc, dupes = cd.compute_dept(wbd, ci, sheet, hr, pf, pt)
        dept_results[sheet] = (res, rl, rc)
        pcs_by_sheet[sheet] = (res, pcs)
    oc_res, oc_pcs = c12_oc.compute_oc(wbd, ci, pf, pt, dept_results=dept_results)
    pcs_by_sheet['c12_oc'] = (oc_res, oc_pcs)

    total = 0
    for sheet, (res, pcs) in pcs_by_sheet.items():
        w = write_o_rows(wb[sheet], res, pcs)
        total += len(w)
        if verbose:
            print(f"  [{sheet:9}] wrote {len(w)} o-cells")
    n = save_clean.save_clean(wb, out_path)
    if verbose:
        print(f"  saved {out_path}  ({total} o-cells written, {n} empty <v> stripped)")
    return total
