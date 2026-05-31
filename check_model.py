import openpyxl
import os
import glob
from openpyxl.utils import get_column_letter

def check_error_matrix(filepath=None):
    """
    Check lgh_diags error matrix (cols X-AG, cols 24-33) for non-zero values.
    Error matrix = live model minus hardwired baseline.
    Returns True if all zeros, False if any differences found.
    """
    if filepath is None:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        files = glob.glob(os.path.join(script_dir, '*.xlsx'))
        if not files:
            print("ERROR: No .xlsx file found in folder.")
            return False
        filepath = files[0]

    print(f"Checking: {os.path.basename(filepath)}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb['lgh_diags']

    # Error matrix: cols X-AG (24-33), rows 4 onwards
    ERROR_COL_START = 24  # X
    ERROR_COL_END   = 33  # AG
    ERROR_ROW_START = 4
    ERROR_ROW_END   = ws.max_row

    # Get period labels from row 3
    period_labels = {}
    for col in range(ERROR_COL_START, ERROR_COL_END + 1):
        v = ws.cell(3, col).value
        period_labels[col] = str(v) if v else f"col{col}"

    errors = []
    for r in range(ERROR_ROW_START, ERROR_ROW_END + 1):
        label = ws.cell(r, 2).value or ws.cell(r, 1).value or ''
        for col in range(ERROR_COL_START, ERROR_COL_END + 1):
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)) and abs(v) > 0.5:
                errors.append((r, col, label, period_labels[col], v))

    if errors:
        print(f"\n❌ ERROR MATRIX: {len(errors)} non-zero cells found:\n")
        print(f"  {'Row':<5} {'Label':<35} {'Period':<12} {'Difference':>15}")
        print("  " + "-" * 70)
        for r, col, label, period, v in errors:
            print(f"  {r:<5} {str(label):<35} {period:<12} {v:>15,.0f}")
        return False
    else:
        print(f"\n✓ ERROR MATRIX: All zeros — model matches baseline.")
        return True

if __name__ == '__main__':
    check_error_matrix()
